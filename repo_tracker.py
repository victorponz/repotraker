#!/usr/bin/env python3
"""
repo_tracker.py

Sigue el progreso de uno o varios repositorios git: por cada commit nuevo
desde la última comprobación, calcula líneas añadidas y eliminadas y las
presenta en una tabla.

No depende de la API de GitHub para listar commits (sin límites de
peticiones): clona cada repositorio en local (como "bare repo") y usa
`git log --numstat`. Para repositorios privados, pasa un token de GitHub
vía --token o la variable de entorno GITHUB_TOKEN; el token se usa solo en
memoria para autenticar el clone/fetch y no se guarda en ningún fichero.

Persistencia:
- Por defecto, el estado (último commit revisado por repo) se guarda en
  un fichero JSON local (~/.repo_tracker/state.json).
- Si configuras Supabase (--supabase-url/--supabase-key o las variables de
  entorno SUPABASE_URL/SUPABASE_KEY), el estado Y el historial completo de
  cada commit (fecha, añadidas, eliminadas) se guardan en tu proyecto de
  Supabase en vez de en el fichero local. Usa la "service_role" key (no la
  "anon"), porque este script corre en tu máquina, no en un navegador, y
  esa key evita líos con Row Level Security. Antes de usarlo, crea las
  tablas ejecutando una vez supabase_schema.sql en el SQL editor de tu
  proyecto.

Requisitos: tener `git` instalado. Solo librería estándar de Python.

Uso rápido (estado local):
    python3 repo_tracker.py add https://github.com/victorponz/javafx
    python3 repo_tracker.py check victorponz/javafx
    python3 repo_tracker.py check --all
    python3 repo_tracker.py list
    python3 repo_tracker.py remove victorponz/javafx

Añadir varios repositorios desde un fichero de texto (uno por línea):
    python3 repo_tracker.py import repos.txt
    # repos.txt admite líneas vacías y comentarios con #
    #   https://github.com/victorponz/javafx
    #   victorponz/otro-repo

Uso con Supabase (estado + historial completo en la nube):
    export SUPABASE_URL=https://xxxx.supabase.co
    export SUPABASE_KEY=eyJ...   # service_role key
    python3 repo_tracker.py add victorponz/javafx
    python3 repo_tracker.py check victorponz/javafx
    python3 repo_tracker.py history victorponz/javafx

Repositorio privado (token solo para esta ejecución):
    export GITHUB_TOKEN=ghp_xxx
    python3 repo_tracker.py check owner/repo-privado

Exportar a Excel (una hoja por repositorio, commits enlazados a GitHub):
    python3 repo_tracker.py xlsx informe.xlsx repos.txt
    # repos.txt tiene el mismo formato que el de 'import': un repositorio
    # por línea (URL o 'owner/repo'), líneas vacías y comentarios con #
    # Requiere openpyxl: pip install openpyxl
"""

import argparse
import base64
import json
import os
import re
import shutil
import subprocess
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

STATE_DIR = Path.home() / ".repo_tracker"
STATE_FILE = STATE_DIR / "state.json"
CACHE_DIR = STATE_DIR / "cache"

REPOS_TABLE = "repo_tracker_repos"
COMMITS_TABLE = "repo_tracker_commits"


# --------------------------------------------------------------------------
# Estado local (fichero JSON) — backend por defecto
# --------------------------------------------------------------------------

def load_local_state() -> dict:
    if STATE_FILE.exists():
        return json.loads(STATE_FILE.read_text())
    return {"repos": {}}


def save_local_state(state: dict) -> None:
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    STATE_FILE.write_text(json.dumps(state, indent=2, ensure_ascii=False))


# --------------------------------------------------------------------------
# API de GitHub: nombre completo del propietario de un repo
# --------------------------------------------------------------------------

def fetch_github_display_name(owner: str, token: str | None) -> str:
    """Consulta la API pública de GitHub para obtener el nombre completo del
    usuario/organización 'owner'. Si no lo tiene puesto en el perfil (campo
    'name' vacío/null) o falla la consulta, devuelve el propio username."""
    url = f"https://api.github.com/users/{urllib.parse.quote(owner)}"
    headers = {
        "Accept": "application/vnd.github+json",
        "User-Agent": "repo_tracker.py",
    }
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(url, headers=headers, method="GET")
    try:
        with urllib.request.urlopen(req) as resp:
            data = json.loads(resp.read())
    except urllib.error.HTTPError as e:
        print(f"Aviso: no se pudo obtener el nombre completo de '{owner}' (HTTP {e.code}); uso el username.")
        return owner
    except urllib.error.URLError as e:
        print(f"Aviso: no se pudo obtener el nombre completo de '{owner}' ({e.reason}); uso el username.")
        return owner
    name = (data.get("name") or "").strip()
    return name or owner


# --------------------------------------------------------------------------
# Cliente REST de Supabase (solo librería estándar, vía urllib)
# --------------------------------------------------------------------------

class SupabaseError(RuntimeError):
    pass


def _sb_call(base_url: str, api_key: str, method: str, table: str,
             params: dict | None = None, body=None, prefer: str | None = None):
    url = f"{base_url}/rest/v1/{table}"
    if params:
        url += "?" + urllib.parse.urlencode(params)
    headers = {
        "apikey": api_key,
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    if prefer:
        headers["Prefer"] = prefer
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as resp:
            raw = resp.read()
            return json.loads(raw) if raw else None
    except urllib.error.HTTPError as e:
        err_body = e.read().decode(errors="replace")
        raise SupabaseError(f"Supabase {method} {table} -> HTTP {e.code}: {err_body}")
    except urllib.error.URLError as e:
        raise SupabaseError(f"No se pudo conectar a Supabase ({base_url}): {e.reason}")


# --------------------------------------------------------------------------
# Store: abstrae el backend de persistencia (local o Supabase)
# --------------------------------------------------------------------------

class Store:
    def __init__(self, supabase_url: str | None, supabase_key: str | None):
        self.supabase_url = supabase_url.rstrip("/") if supabase_url else None
        self.supabase_key = supabase_key
        self.use_supabase = bool(self.supabase_url and self.supabase_key)
        self._local = None if self.use_supabase else load_local_state()

    def _call(self, method, table, params=None, body=None, prefer=None):
        return _sb_call(self.supabase_url, self.supabase_key, method, table, params, body, prefer)

    # ---- repos ----

    def get_repo(self, key: str) -> dict | None:
        if self.use_supabase:
            rows = self._call("GET", REPOS_TABLE, params={"key": f"eq.{key}", "select": "*"})
            return rows[0] if rows else None
        return self._local["repos"].get(key)

    def list_repo_keys(self) -> list[str]:
        if self.use_supabase:
            rows = self._call("GET", REPOS_TABLE, params={"select": "key"})
            return [r["key"] for r in rows]
        return list(self._local["repos"].keys())

    def upsert_repo(self, key: str, row: dict) -> None:
        if self.use_supabase:
            payload = {"key": key, **row}
            self._call(
                "POST", REPOS_TABLE,
                params={"on_conflict": "key"},
                body=[payload],
                prefer="resolution=merge-duplicates,return=minimal",
            )
        else:
            self._local["repos"][key] = row

    def delete_repo(self, key: str) -> None:
        if self.use_supabase:
            self._call("DELETE", REPOS_TABLE, params={"key": f"eq.{key}"})
        else:
            self._local["repos"].pop(key, None)

    # ---- historial de commits (solo tiene efecto real con Supabase) ----

    def record_commits(self, key: str, commits: list[dict]) -> None:
        if not self.use_supabase or not commits:
            return
        payload = [{
            "repo_key": key,
            "num": c["num"],
            "sha": c["sha"],
            "commit_date": c["date"],
            "subject": c["subject"],
            "added": c["added"],
            "removed": c["removed"],
        } for c in commits]
        self._call(
            "POST", COMMITS_TABLE,
            params={"on_conflict": "repo_key,sha"},
            body=payload,
            prefer="resolution=merge-duplicates,return=minimal",
        )

    def get_commit_history(self, key: str) -> list[dict]:
        if not self.use_supabase:
            return []
        rows = self._call(
            "GET", COMMITS_TABLE,
            params={"repo_key": f"eq.{key}", "select": "*", "order": "num.asc"},
        )
        return [{
            "num": r["num"], "sha": r["sha"], "date": r["commit_date"],
            "subject": r["subject"] or "", "added": r["added"], "removed": r["removed"],
        } for r in rows]

    # ---- persistencia final ----

    def flush(self) -> None:
        if not self.use_supabase:
            save_local_state(self._local)


def resolve_supabase_config(args) -> tuple[str | None, str | None]:
    url = getattr(args, "supabase_url", None) or os.environ.get("SUPABASE_URL")
    key = getattr(args, "supabase_key", None) or os.environ.get("SUPABASE_KEY")
    return url, key


def make_store(args) -> Store:
    url, key = resolve_supabase_config(args)
    if bool(url) != bool(key):
        print("Aviso: para usar Supabase necesitas tanto la URL como la key; usando almacenamiento local.")
        url, key = None, None
    return Store(url, key)


# --------------------------------------------------------------------------
# Normalización de la referencia al repositorio
# --------------------------------------------------------------------------

def normalize_repo(ref: str) -> tuple[str, str]:
    """Devuelve (clave 'owner/repo', url https) a partir de una URL o de 'owner/repo'."""
    ref = ref.strip().rstrip("/")
    ref = re.sub(r"\.git$", "", ref)

    m = re.match(r"(?:https?://github\.com/|git@github\.com:)([^/]+)/([^/]+)$", ref)
    if m:
        owner, repo = m.group(1), m.group(2)
        return f"{owner}/{repo}", f"https://github.com/{owner}/{repo}.git"

    m = re.match(r"^([^/\s]+)/([^/\s]+)$", ref)
    if m:
        owner, repo = m.group(1), m.group(2)
        return f"{owner}/{repo}", f"https://github.com/{owner}/{repo}.git"

    raise ValueError(
        f"No reconozco '{ref}' como repositorio. Usa una URL de GitHub o 'owner/repo'."
    )


def cache_path(key: str) -> Path:
    return CACHE_DIR / key.replace("/", "__")


def resolve_token(args) -> str | None:
    """El --token explícito gana; si no, se usa la variable de entorno GITHUB_TOKEN."""
    token = getattr(args, "token", None)
    return token or os.environ.get("GITHUB_TOKEN")


def git_auth_opts(token: str | None) -> list[str]:
    """Opciones globales de git para autenticar por HTTPS con un token,
    sin que el token quede grabado en el remote ni en ningún fichero de
    configuración: se manda como cabecera HTTP solo para esta invocación."""
    if not token:
        return []
    raw = f"x-access-token:{token}".encode()
    header = "Authorization: Basic " + base64.b64encode(raw).decode()
    return ["-c", f"http.extraHeader={header}"]


# --------------------------------------------------------------------------
# Operaciones git
# --------------------------------------------------------------------------

def run_git(args: list[str]) -> str:
    env = {**os.environ, "GIT_TERMINAL_PROMPT": "0"}
    result = subprocess.run(["git", *args], capture_output=True, text=True, env=env)
    if result.returncode != 0:
        stderr = result.stderr.strip()
        if "could not read Username" in stderr or "Authentication failed" in stderr or "terminal prompts disabled" in stderr:
            stderr += "\n  -> credenciales rechazadas: revisa que el token sea correcto y tenga permiso de lectura sobre el repo."
        raise RuntimeError(f"git {' '.join(args)} -> {stderr}")
    return result.stdout


def clone_bare(url: str, path: Path, token: str | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    run_git([*git_auth_opts(token), "clone", "--quiet", "--bare", url, str(path)])


def is_valid_bare_repo(path: Path) -> bool:
    """Comprueba si 'path' ya contiene un clon bare de git utilizable."""
    if not path.is_dir():
        return False
    result = subprocess.run(
        ["git", "--git-dir", str(path), "rev-parse", "--is-bare-repository"],
        capture_output=True, text=True,
    )
    return result.returncode == 0 and result.stdout.strip() == "true"


def ensure_cloned(url: str, path: Path, token: str | None = None) -> bool:
    """Se asegura de que 'path' contiene un clon bare válido de 'url'.
    Si ya existe uno, no hace nada (idempotente) y devuelve False.
    Si 'path' existe pero está corrupto/incompleto (por ejemplo, de una
    ejecución anterior interrumpida a medias), lo borra y clona de nuevo.
    Devuelve True si se ha clonado de verdad."""
    if path.exists():
        if is_valid_bare_repo(path):
            return False
        shutil.rmtree(path)
    clone_bare(url, path, token)
    return True


def fetch(path: Path, token: str | None = None) -> None:
    # En un bare clone las ramas viven en refs/heads/*, así que pedimos
    # explícitamente que el fetch las actualice ahí (modo espejo).
    run_git([
        *git_auth_opts(token),
        "--git-dir", str(path), "fetch", "--quiet", "--prune",
        "origin", "+refs/heads/*:refs/heads/*",
    ])


def default_branch(path: Path) -> str:
    out = run_git(["--git-dir", str(path), "symbolic-ref", "HEAD"])
    return out.strip().split("/")[-1]


def rev_parse(path: Path, branch: str) -> str:
    return run_git(["--git-dir", str(path), "rev-parse", f"refs/heads/{branch}"]).strip()


COMMIT_MARK = "\x01COMMIT\x01"


def log_numstat(path: Path, rev_range: str) -> str:
    return run_git([
        "--git-dir", str(path),
        "log", rev_range,
        "--reverse",
        "--numstat",
        f"--pretty=format:{COMMIT_MARK}%H\x02%ad\x02%s",
        "--date=iso-strict",
    ])


def parse_log(raw: str) -> list[dict]:
    commits = []
    blocks = raw.split(COMMIT_MARK)
    for block in blocks:
        block = block.strip("\n")
        if not block:
            continue
        header, _, body = block.partition("\n")
        sha, date, subject = header.split("\x02", 2)
        added = removed = 0
        for line in body.splitlines():
            line = line.strip()
            if not line:
                continue
            parts = line.split("\t")
            if len(parts) != 3:
                continue
            a, r, _file = parts
            if a.isdigit():
                added += int(a)
            if r.isdigit():
                removed += int(r)
        commits.append({
            "sha": sha,
            "date": date,
            "subject": subject,
            "added": added,
            "removed": removed,
        })
    return commits


# --------------------------------------------------------------------------
# Tabla de salida
# --------------------------------------------------------------------------

def print_table(rows: list[dict]) -> None:
    headers = ["#", "Commit", "Fecha", "Añadidas", "Eliminadas", "Mensaje"]
    data = []
    for r in rows:
        data.append([
            str(r["num"]), r["sha"][:7], r["date"][:19],
            f"+{r['added']}", f"-{r['removed']}", r["subject"][:40],
        ])
    widths = [max(len(h), *(len(row[i]) for row in data)) if data else len(h)
              for i, h in enumerate(headers)]
    line = "  ".join(h.ljust(widths[i]) for i, h in enumerate(headers))
    print(line)
    print("  ".join("-" * w for w in widths))
    for row in data:
        print("  ".join(cell.ljust(widths[i]) for i, cell in enumerate(row)))


def write_csv(rows: list[dict], path: Path) -> None:
    import csv
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["numero", "commit", "fecha", "añadidas", "eliminadas", "mensaje"])
        for r in rows:
            writer.writerow([r["num"], r["sha"], r["date"], r["added"], r["removed"], r["subject"]])


# --------------------------------------------------------------------------
# Comandos
# --------------------------------------------------------------------------

def _add_one(store: Store, ref: str, branch_override: str | None, token: str | None) -> str | None:
    """Añade un único repositorio a 'store'. Devuelve la key si queda en seguimiento
    (recién añadido o ya lo estaba), o None si falló el clonado.
    Lanza ValueError si la referencia no se reconoce."""
    key, url = normalize_repo(ref)

    if store.get_repo(key):
        print(f"'{key}' ya se está siguiendo.")
        return key

    path = cache_path(key)
    try:
        cloned = ensure_cloned(url, path, token)
    except RuntimeError as e:
        hint = "" if token else " (si es un repositorio privado, pasa --token o define GITHUB_TOKEN)"
        print(f"Error al clonar{hint}: {e}")
        return None
    if cloned:
        print(f"Clonado {url}.")
    else:
        print(f"Ya existe un clon local de '{key}', lo reutilizo.")

    branch = branch_override or default_branch(path)
    store.upsert_repo(key, {
        "url": url,
        "branch": branch,
        "last_sha": None,
        "commit_count": 0,
        "total_added": 0,
        "total_removed": 0,
        "last_checked": None,
    })
    backend = "Supabase" if store.use_supabase else "fichero local"
    print(f"'{key}' añadido (rama: {branch}, estado en {backend}).")
    return key


def cmd_add(args):
    token = resolve_token(args)
    store = make_store(args)
    key = _add_one(store, args.repo, args.branch, token)
    if key:
        print(f"Ejecuta 'check {key}' para ver el historial.")
    store.flush()


def read_repo_file(path: Path) -> list[str]:
    """Lee un fichero de texto con un repositorio por línea.
    Ignora líneas vacías y líneas que empiecen por '#' (comentarios)."""
    if not path.exists():
        raise FileNotFoundError(f"No existe el fichero '{path}'.")
    refs = []
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        refs.append(line)
    return refs


def cmd_import(args):
    path = Path(args.file)
    try:
        refs = read_repo_file(path)
    except FileNotFoundError as e:
        print(str(e))
        return

    if not refs:
        print(f"'{path}' no contiene ningún repositorio.")
        return

    token = resolve_token(args)
    store = make_store(args)
    auth_hint = "" if token else " (si es privado, pasa --token o define GITHUB_TOKEN)"
    print(f"Importando {len(refs)} repositorio(s) desde '{path}'...")
    for ref in refs:
        try:
            key = _add_one(store, ref, args.branch, token)
        except ValueError as e:
            print(f"Aviso: {e}")
            continue
        if key is None:
            continue
        try:
            _check_one(store, key, args.csv, token)
        except RuntimeError as e:
            print(f"[{key}] Error{auth_hint}: {e}")
    store.flush()


def _commits_for_export(store: Store, key: str, info: dict) -> list[dict] | None:
    """Devuelve el historial completo de commits de un repo para exportarlo,
    con 'num' asignado en cada commit. Usa el histórico guardado en Supabase
    si existe; si no (backend local, o todavía no se ha hecho ningún check),
    lo reconstruye leyendo el log completo del clon local en caché.
    Devuelve None si no hay ningún dato disponible (ni Supabase ni clon local)."""
    if store.use_supabase:
        commits = store.get_commit_history(key)
        if commits:
            return commits

    path = cache_path(key)
    if not (path.exists() and is_valid_bare_repo(path)):
        return None
    branch = info.get("branch") or default_branch(path)
    raw = log_numstat(path, branch)
    commits = parse_log(raw)
    for i, c in enumerate(commits, start=1):
        c["num"] = i
    return commits


def _sheet_name_factory():
    """Genera nombres de hoja válidos y únicos para Excel a partir del
    nombre completo del propietario del repo, desambiguando con el nombre
    del repo si dos coincidieran tras limpiar caracteres no válidos."""
    used: dict[str, bool] = {}
    invalid_chars = set('[]:*?/\\')

    def clean(s: str) -> str:
        s = "".join(c for c in s if c not in invalid_chars).strip()
        return (s or "repo")[:31]

    def make(display_name: str, repo_name: str) -> str:
        candidate = clean(display_name)
        if candidate.lower() in used:
            candidate = clean(f"{display_name} ({repo_name})")
        n = 2
        while candidate.lower() in used:
            suffix = f" ({n})"
            candidate = clean(display_name)[: 31 - len(suffix)] + suffix
            n += 1
        used[candidate.lower()] = True
        return candidate

    return make


def cmd_xlsx(args):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("Esta función necesita la librería 'openpyxl'. Instálala con: pip install openpyxl")
        return

    store = make_store(args)
    token = resolve_token(args)

    repos_path = Path(args.repos_file)
    try:
        refs = read_repo_file(repos_path)
    except FileNotFoundError as e:
        print(str(e))
        return
    if not refs:
        print(f"'{repos_path}' no contiene ningún repositorio.")
        return

    keys = []
    for ref in refs:
        try:
            key, _ = normalize_repo(ref)
        except ValueError as e:
            print(f"Aviso: {e}")
            continue
        keys.append(key)

    if not keys:
        print(f"'{repos_path}' no contiene ningún repositorio válido.")
        return

    wb = Workbook()
    wb.remove(wb.active)
    sheet_name_for = _sheet_name_factory()

    header_font = Font(name="Arial", bold=True)
    normal_font = Font(name="Arial")
    link_font = Font(name="Arial", color="0563C1", underline="single")
    headers = ["#", "Commit", "Fecha", "Mensaje", "Añadidas", "Eliminadas"]
    col_widths = [6, 12, 22, 70, 10, 10]

    exported = 0
    display_name_cache: dict[str, str] = {}
    for key in keys:
        info = store.get_repo(key) or {"branch": None}
        owner, repo_name = key.split("/", 1)
        if owner not in display_name_cache:
            display_name_cache[owner] = fetch_github_display_name(owner, token)
        display_name = display_name_cache[owner]
        commits = _commits_for_export(store, key, info)

        ws = wb.create_sheet(title=sheet_name_for(display_name, repo_name))
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
        ws.freeze_panes = "A2"

        if commits is None:
            ws.append(["", "(sin clon local; ejecuta 'check' o 'import' primero)", "", "", "", ""])
            for cell in ws[2]:
                cell.font = normal_font
        else:
            for c in commits:
                ws.append([c.get("num", ""), c["sha"][:7], c["date"], c["subject"], c["added"], c["removed"]])
                row_idx = ws.max_row
                for col in range(1, 7):
                    ws.cell(row=row_idx, column=col).font = normal_font
                commit_cell = ws.cell(row=row_idx, column=2)
                commit_cell.hyperlink = f"https://github.com/{owner}/{repo_name}/commit/{c['sha']}"
                commit_cell.font = link_font
            exported += 1

        for i, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(args.file)
    print(f"Generado '{args.file}' con {len(keys)} hoja(s) ({exported} con datos de commits).")


def _check_one(store: Store, key: str, export_csv: bool, token: str | None) -> None:
    info = store.get_repo(key)
    path = cache_path(key)
    if path.exists() and is_valid_bare_repo(path):
        fetch(path, token)
    else:
        print(f"[{key}] No hay clon local válido, lo creo...")
        ensure_cloned(info["url"], path, token)

    branch = info.get("branch") or default_branch(path)
    head = rev_parse(path, branch)

    last_sha = info.get("last_sha")
    if last_sha is None:
        rev_range = head  # todo el historial hasta head
    elif last_sha == head:
        print(f"[{key}] Sin commits nuevos desde la última comprobación.")
        info["last_checked"] = datetime.now(timezone.utc).isoformat()
        store.upsert_repo(key, info)
        return
    else:
        rev_range = f"{last_sha}..{head}"

    raw = log_numstat(path, rev_range)
    commits = parse_log(raw)

    if not commits:
        print(f"[{key}] Sin commits nuevos desde la última comprobación.")
        info["last_sha"] = head
        info["last_checked"] = datetime.now(timezone.utc).isoformat()
        store.upsert_repo(key, info)
        return

    start = info.get("commit_count", 0)
    for i, c in enumerate(commits, start=1):
        c["num"] = start + i

    print(f"\n=== {key} (rama: {branch}) — {len(commits)} commit(s) nuevo(s) ===")
    print_table(commits)
    total_added = sum(c["added"] for c in commits)
    total_removed = sum(c["removed"] for c in commits)
    print(f"\nTotal añadidas: +{total_added}   Total eliminadas: -{total_removed}")

    if export_csv:
        csv_path = Path(f"{key.replace('/', '__')}_commits.csv")
        write_csv(commits, csv_path)
        print(f"Guardado en {csv_path}")

    store.record_commits(key, commits)

    info["last_sha"] = head
    info["commit_count"] = start + len(commits)
    info["total_added"] = info.get("total_added", 0) + total_added
    info["total_removed"] = info.get("total_removed", 0) + total_removed
    info["last_checked"] = datetime.now(timezone.utc).isoformat()
    store.upsert_repo(key, info)


def cmd_check(args):
    store = make_store(args)
    token = resolve_token(args)
    auth_hint = "" if token else " (si es privado, pasa --token o define GITHUB_TOKEN)"

    if args.all:
        keys = store.list_repo_keys()
        if not keys:
            print("No hay repositorios en seguimiento. Usa 'add' primero.")
            return
        for key in keys:
            try:
                _check_one(store, key, args.csv, token)
            except RuntimeError as e:
                print(f"[{key}] Error{auth_hint}: {e}")
        store.flush()
        return

    if not args.repo:
        print("Indica un repositorio o usa --all.")
        return
    key, _ = normalize_repo(args.repo)
    if not store.get_repo(key):
        print(f"'{key}' no está en seguimiento. Usa 'add {args.repo}' primero.")
        return
    try:
        _check_one(store, key, args.csv, token)
    except RuntimeError as e:
        print(f"[{key}] Error{auth_hint}: {e}")
    store.flush()


def cmd_list(args):
    store = make_store(args)
    keys = store.list_repo_keys()
    if not keys:
        print("No hay repositorios en seguimiento.")
        return
    for key in keys:
        info = store.get_repo(key) or {}
        last = info.get("last_checked") or "nunca"
        totals = f"+{info.get('total_added', 0)}/-{info.get('total_removed', 0)}" if info.get("commit_count") else ""
        print(f"{key}  (rama: {info.get('branch')}, commits: {info.get('commit_count', 0)} {totals}, última comprobación: {last})")


def cmd_remove(args):
    key, _ = normalize_repo(args.repo)
    store = make_store(args)
    if not store.get_repo(key):
        print(f"'{key}' no estaba en seguimiento.")
        return
    store.delete_repo(key)
    store.flush()
    extra = "" if store.use_supabase else " (el clon local en cache se conserva, bórralo a mano si quieres liberar espacio)"
    print(f"'{key}' eliminado del seguimiento{extra}.")


def cmd_history(args):
    store = make_store(args)
    if not store.use_supabase:
        print("El historial completo de commits solo se guarda cuando usas Supabase "
              "(define --supabase-url/--supabase-key o SUPABASE_URL/SUPABASE_KEY).")
        return
    key, _ = normalize_repo(args.repo)
    if not store.get_repo(key):
        print(f"'{key}' no está en seguimiento.")
        return
    commits = store.get_commit_history(key)
    if not commits:
        print(f"Todavía no hay commits guardados para '{key}'. Ejecuta 'check {key}' primero.")
        return
    print(f"\n=== Historial completo de {key} — {len(commits)} commit(s) ===")
    print_table(commits)
    if args.csv:
        csv_path = Path(f"{key.replace('/', '__')}_historial.csv")
        write_csv(commits, csv_path)
        print(f"Guardado en {csv_path}")


def main():
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--supabase-url", help="URL del proyecto Supabase. También: variable de entorno SUPABASE_URL.")
    common.add_argument("--supabase-key", help="Service role key de Supabase. También: variable de entorno SUPABASE_KEY.")

    parser = argparse.ArgumentParser(description="Seguimiento de commits en repositorios git.")
    sub = parser.add_subparsers(dest="command", required=True)

    p_add = sub.add_parser("add", help="Empieza a seguir un repositorio", parents=[common])
    p_add.add_argument("repo", help="URL de GitHub u 'owner/repo'")
    p_add.add_argument("--branch", help="Rama a seguir (por defecto, la rama por defecto del repo)")
    p_add.add_argument("--token", help="Token de GitHub (para repos privados). También se puede definir GITHUB_TOKEN.")
    p_add.set_defaults(func=cmd_add)

    p_import = sub.add_parser("import", help="Añade varios repositorios leyéndolos de un fichero de texto (uno por línea)", parents=[common])
    p_import.add_argument("file", help="Ruta al fichero de texto con un repositorio por línea")
    p_import.add_argument("--branch", help="Rama a seguir para todos los repositorios (por defecto, la rama por defecto de cada uno)")
    p_import.add_argument("--token", help="Token de GitHub (para repos privados). También se puede definir GITHUB_TOKEN.")
    p_import.add_argument("--csv", action="store_true", help="Exporta también un CSV con los resultados de cada check")
    p_import.set_defaults(func=cmd_import)

    p_xlsx = sub.add_parser("xlsx", help="Genera un Excel con una hoja por repositorio y enlaces a cada commit en GitHub", parents=[common])
    p_xlsx.add_argument("file", help="Ruta del fichero .xlsx a generar")
    p_xlsx.add_argument("repos_file", help="Fichero de texto con un repositorio por línea (mismo formato que 'import')")
    p_xlsx.set_defaults(func=cmd_xlsx)

    p_check = sub.add_parser("check", help="Comprueba commits nuevos desde la última vez", parents=[common])
    p_check.add_argument("repo", nargs="?", help="URL de GitHub u 'owner/repo'")
    p_check.add_argument("--all", action="store_true", help="Comprueba todos los repositorios seguidos")
    p_check.add_argument("--csv", action="store_true", help="Exporta también un CSV con los resultados")
    p_check.add_argument("--token", help="Token de GitHub (para repos privados). También se puede definir GITHUB_TOKEN.")
    p_check.set_defaults(func=cmd_check)

    p_list = sub.add_parser("list", help="Lista los repositorios en seguimiento", parents=[common])
    p_list.set_defaults(func=cmd_list)

    p_remove = sub.add_parser("remove", help="Deja de seguir un repositorio", parents=[common])
    p_remove.add_argument("repo", help="URL de GitHub u 'owner/repo'")
    p_remove.set_defaults(func=cmd_remove)

    p_history = sub.add_parser("history", help="Muestra el historial completo guardado en Supabase", parents=[common])
    p_history.add_argument("repo", help="URL de GitHub u 'owner/repo'")
    p_history.add_argument("--csv", action="store_true", help="Exporta también un CSV con el historial")
    p_history.set_defaults(func=cmd_history)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
